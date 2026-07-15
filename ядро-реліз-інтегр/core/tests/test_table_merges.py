import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document

from core.normalize_docx import normalize_docx


def _merge_counts(path: Path) -> tuple[int, int]:
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    return xml.count(b"<w:gridSpan"), xml.count(b"<w:vMerge")


class TableMergeIntegrationTests(unittest.TestCase):
    def test_normalize_preserves_horizontal_and_vertical_merges(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "source.docx"
            output = root / "output.docx"

            doc = Document()
            table = doc.add_table(rows=3, cols=3)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "horizontal"
            table.cell(1, 2).merge(table.cell(2, 2)).text = "vertical"
            doc.save(source)

            normalize_docx(source, output)

            source_counts = _merge_counts(source)
            output_counts = _merge_counts(output)
            self.assertGreater(source_counts[0], 0)
            self.assertGreater(source_counts[1], 0)
            self.assertEqual(output_counts, source_counts)


if __name__ == "__main__":
    unittest.main()
